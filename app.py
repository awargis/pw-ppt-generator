"""
PW Test PPT Generator — Automated Question Paper → Subject-wise PPT Discussion Decks
=====================================================================================

This app converts a scanned question paper PDF into subject-wise PPT discussion decks.
It tries to detect each question as a separate crop, assign a subject using section
headers, and generate one PPT per subject using a sample template slide.

Requirements:
- pip install streamlit pdf2image pillow python-pptx google-genai numpy
- poppler-utils installed on the system for pdf2image

Run:
- streamlit run app.py
"""

import copy
import io
import json
import os
import re
import shutil
import zipfile
from dataclasses import dataclass
from typing import Optional

import numpy as np
import streamlit as st
from PIL import Image, ImageEnhance, ImageOps
from pdf2image import convert_from_bytes
from pptx import Presentation
from pptx.util import Emu

from google import genai
from google.genai import types


# ---------------------------------------------------------------------------
# Streamlit page setup
# ---------------------------------------------------------------------------
st.set_page_config(page_title="PW Test PPT Generator", layout="wide")
st.title("📚 Automated Question Paper → PPT Generator")
st.caption(
    "Upload the question paper PDF + your fixed sample PPT once. "
    "Get perfectly cropped, subject-wise, one-question-per-slide decks — every time."
)


# ---------------------------------------------------------------------------
# Sidebar config
# ---------------------------------------------------------------------------
st.sidebar.header("⚙️ Configuration")
gemini_api_key = st.sidebar.text_input("Gemini API Key", type="password")
MODEL_CHOICES = [
    "gemini-2.5-flash",
    "gemini-2.5-pro",
    "gemini-2.0-flash",
]
model_name = st.sidebar.selectbox(
    "Gemini model",
    MODEL_CHOICES,
    help="Use Flash for speed or Pro for better OCR/vision accuracy on dense pages.",
)

exam_type = st.sidebar.selectbox("Exam Type", ["JEE Main / Advanced", "NEET"])
KNOWN_SUBJECTS = (
    ["Physics", "Chemistry", "Mathematics"]
    if "JEE" in exam_type
    else ["Physics", "Chemistry", "Botany", "Zoology"]
)

st.sidebar.info(f"Expected subjects: **{', '.join(KNOWN_SUBJECTS)}**")

with st.sidebar.expander("Advanced (crop tuning)"):
    render_dpi = st.slider("PDF render DPI", 200, 400, 300, step=50)
    pad_px = st.slider("Padding around each question (px @ render DPI)", 5, 40, 18)
    col_gap_pct = st.slider("Column-gap exclusion (%)", 0.5, 4.0, 1.5, step=0.5)
    outer_margin_pct = st.slider("Outer page-margin exclusion (%)", 0.0, 3.0, 1.2, step=0.2)
    invert_question_images = st.checkbox("Invert question images to black background", value=False)


# ---------------------------------------------------------------------------
# Data models
# ---------------------------------------------------------------------------
@dataclass
class DetectedItem:
    kind: str
    page_index: int
    column_index: int
    y0: int
    x0: int
    y1: int
    x1: int
    question_number: Optional[int] = None
    subject: Optional[str] = None
    confidence: float = 0.0

    @property
    def sort_key(self):
        return (self.page_index, self.column_index, self.y0)


# ---------------------------------------------------------------------------
# Subject name normalization helpers
# ---------------------------------------------------------------------------
SUBJECT_ALIASES = {
    "physics": "Physics",
    "phy": "Physics",
    "chemistry": "Chemistry",
    "chem": "Chemistry",
    "mathematics": "Mathematics",
    "math": "Mathematics",
    "maths": "Mathematics",
    "botany": "Botany",
    "zoology": "Zoology",
}


def normalize_subject(value: Optional[str], known_subjects: list[str]) -> Optional[str]:
    if value is None:
        return None

    cleaned = str(value).strip()
    if not cleaned:
        return None

    cleaned = cleaned.lower()
    cleaned = cleaned.replace("section", " ")
    cleaned = re.sub(r"[^a-z\s]", " ", cleaned)
    cleaned = re.sub(r"\s+", " ", cleaned).strip()

    if not cleaned:
        return None

    if cleaned in SUBJECT_ALIASES:
        candidate = SUBJECT_ALIASES[cleaned]
        if candidate in known_subjects:
            return candidate

    for subject in known_subjects:
        subject_key = subject.lower()
        if subject_key in cleaned or cleaned in subject_key:
            return subject

    return None


# ---------------------------------------------------------------------------
# Gemini helpers
# ---------------------------------------------------------------------------
DETECTION_PROMPT = """You are analysing ONE COLUMN of a scanned exam question-paper page.

Return a JSON array. Each element is ONE of the following two shapes:

1) A question:
   {
     "type": "question",
     "question_number": <integer>,
     "box_2d": [ymin, xmin, ymax, xmax],
     "confidence": <number between 0 and 1>,
     "question_format": "mcq|numerical|matching|assertion_reason|passage|descriptive|unknown"
   }
   The box must tightly enclose the ENTIRE question: its number, full text,
   all diagram/figure/table content, and ALL answer options or text, but must stop
   before the next question's number begins. Do not include any part of a previous
   or next question.

2) A section header banner:
   {
     "type": "section_header",
     "subject": "<one of: {subjects}>",
     "box_2d": [ymin, xmin, ymax, xmax],
     "confidence": <number between 0 and 1>
   }

Coordinates are integers normalized 0-1000 relative to THIS image's width/height,
in [ymin, xmin, ymax, xmax] order.
Return ONLY the JSON array, no commentary.
"""

ANSWER_KEY_PROMPT = """This image is an answer key for an exam. Extract every
question number and its correct option/answer. Return ONLY a JSON object
mapping question number (as a string) to the answer, e.g.:
{"1": "3", "2": "1", "3": "4"}
If the answer is a numeric-value (integer type) answer, keep it as given.
"""


def get_client(api_key: str):
    return genai.Client(api_key=api_key)


def _extract_json(text: str):
    text = str(text).strip()
    text = re.sub(r"^```(json)?", "", text).strip()
    text = re.sub(r"```$", "", text).strip()
    if not text:
        raise ValueError("Empty Gemini response")
    return json.loads(text)


def detect_items_in_column(client, model_name, pil_img, subjects) -> list:
    try:
        buf = io.BytesIO()
        pil_img.save(buf, format="PNG")
        prompt = DETECTION_PROMPT.format(subjects=", ".join(subjects))
        resp = client.models.generate_content(
            model=model_name,
            contents=[
                types.Part.from_bytes(data=buf.getvalue(), mime_type="image/png"),
                prompt,
            ],
            config=types.GenerateContentConfig(
                temperature=0,
                response_mime_type="application/json",
            ),
        )
        return _extract_json(resp.text)
    except Exception:
        return []


def ocr_answer_key_image(client, model_name, pil_img) -> dict:
    try:
        buf = io.BytesIO()
        pil_img.save(buf, format="PNG")
        resp = client.models.generate_content(
            model=model_name,
            contents=[
                types.Part.from_bytes(data=buf.getvalue(), mime_type="image/png"),
                ANSWER_KEY_PROMPT,
            ],
            config=types.GenerateContentConfig(
                temperature=0,
                response_mime_type="application/json",
            ),
        )
        raw = _extract_json(resp.text)
        return {int(k): str(v).upper() for k, v in raw.items()}
    except Exception:
        return {}


def parse_answer_key_text(text: str) -> dict:
    """Robust parser for pasted answer keys: '1: (3)', '1) 3', '1-3', '1.(3)' etc."""
    ans_dict = {}
    if not text:
        return ans_dict
    matches = re.findall(r"(\d{1,3})\s*[\.\):\-]+\s*\(?\s*([1-4A-Da-d]+)\s*\)?", text)
    for q_num, ans in matches:
        ans_dict[int(q_num)] = ans.upper()
    return ans_dict


# ---------------------------------------------------------------------------
# Image helpers
# ---------------------------------------------------------------------------
def box_2d_to_abs(box_2d, img_w, img_h, offset_x, offset_y):
    """Convert normalized 0-1000 Gemini box to absolute pixel coords on full page."""
    ymin, xmin, ymax, xmax = box_2d
    x0 = offset_x + int(xmin / 1000 * img_w)
    x1 = offset_x + int(xmax / 1000 * img_w)
    y0 = offset_y + int(ymin / 1000 * img_h)
    y1 = offset_y + int(ymax / 1000 * img_h)
    return x0, y0, x1, y1


def clip_against_neighbours(items_in_column, gap_px=6):
    """Stop adjacent question crops from overlapping each other."""
    items_in_column = sorted(items_in_column, key=lambda it: it.y0)
    for i, item in enumerate(items_in_column):
        if i > 0:
            prev = items_in_column[i - 1]
            floor = prev.y1 + gap_px
            if item.y0 < floor:
                item.y0 = floor
        if i < len(items_in_column) - 1:
            nxt = items_in_column[i + 1]
            ceiling = nxt.y0 - gap_px
            if item.y1 > ceiling:
                item.y1 = ceiling
    return items_in_column


def autotrim(pil_img: Image.Image, pad: int = 18, bg_thresh: int = 245) -> Image.Image:
    """Trim whitespace while retaining a small fixed padding."""
    if pil_img.width < 4 or pil_img.height < 4:
        return pil_img

    gray = pil_img.convert("L")
    arr = np.array(gray)
    mask = arr < bg_thresh
    ys, xs = np.where(mask)
    if len(xs) == 0:
        return pil_img

    x0, x1 = int(xs.min()), int(xs.max())
    y0, y1 = int(ys.min()), int(ys.max())
    x0 = max(0, x0 - pad)
    y0 = max(0, y0 - pad)
    x1 = min(arr.shape[1] - 1, x1 + pad)
    y1 = min(arr.shape[0] - 1, y1 + pad)

    return pil_img.crop((x0, y0, x1, y1))


def invert_to_black_bg(pil_img: Image.Image) -> Image.Image:
    rgb_img = pil_img.convert("RGB")
    inverted_img = ImageOps.invert(rgb_img)
    return ImageEnhance.Contrast(inverted_img).enhance(1.15)


def calculate_crop_size(img_w, img_h, content_w: int, content_h: int, mode: str = "contain"):
    if img_w <= 0 or img_h <= 0:
        return 0, 0
    if mode == "cover":
        scale = max(content_w / img_w, content_h / img_h)
    else:
        scale = min(content_w / img_w, content_h / img_h)
    return max(1, int(img_w * scale)), max(1, int(img_h * scale))


def validate_crop(img: Image.Image, q_number: Optional[int] = None):
    if img is None:
        raise ValueError(f"Question crop is empty for Q{q_number}.")
    if img.width < 80 or img.height < 40:
        raise ValueError(f"Question crop for Q{q_number} is too small ({img.width}x{img.height}).")
    if img.width / img.height > 12 or img.width / img.height < 0.08:
        raise ValueError(f"Question crop for Q{q_number} has suspicious dimensions ({img.width}x{img.height}).")


def split_columns(page_img: Image.Image, outer_margin_pct: float, col_gap_pct: float):
    w, h = page_img.size
    om = outer_margin_pct / 100.0
    cg = col_gap_pct / 100.0

    left_x0 = int(w * om)
    left_x1 = int(w * (0.5 - cg))
    right_x0 = int(w * (0.5 + cg))
    right_x1 = int(w * (1 - om))
    top = int(h * om)
    bottom = int(h * (1 - om))

    left_col = page_img.crop((left_x0, top, left_x1, bottom))
    right_col = page_img.crop((right_x0, top, right_x1, bottom))

    return [(left_col, left_x0, top), (right_col, right_x0, top)]


def deduplicate_questions(items: list[DetectedItem]) -> list[DetectedItem]:
    seen = set()
    result = []
    for item in sorted(items, key=lambda x: x.sort_key):
        key = (item.question_number, item.page_index, item.column_index, round(item.y0), round(item.x0))
        if item.question_number is None:
            result.append(item)
            continue
        if key in seen:
            continue
        seen.add(key)
        result.append(item)
    return result


def remove_overlapping_shapes(slide, left, top, right, bottom):
    """Remove template placeholders or old picture shapes in the content area."""
    to_remove = []
    for shp in slide.shapes:
        try:
            if shp.has_text_frame:
                text = shp.text_frame.text or ""
                if "#QUESTION" in text or "Ans." in text:
                    continue
        except Exception:
            pass

        try:
            shp_left = shp.left
            shp_top = shp.top
            shp_right = shp.left + shp.width
            shp_bottom = shp.top + shp.height
            overlaps = (
                shp_left < right and shp_right > left and shp_top < bottom and shp_bottom > top
            )
            if overlaps:
                to_remove.append(shp)
        except Exception:
            continue

    for shp in to_remove:
        try:
            shp._element.getparent().remove(shp._element)
        except Exception:
            pass


# ---------------------------------------------------------------------------
# PPTX helpers
# ---------------------------------------------------------------------------
_R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def duplicate_slide(prs: Presentation, index: int):
    src = prs.slides[index]
    layout = src.slide_layout
    new_slide = prs.slides.add_slide(layout)

    for shp in list(new_slide.shapes):
        try:
            shp._element.getparent().remove(shp._element)
        except Exception:
            pass

    rid_map = {}
    for old_rId, rel in src.part.rels.items():
        if "notesSlide" in rel.reltype or "slideLayout" in rel.reltype:
            continue
        if rel.is_external:
            new_rId = new_slide.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        else:
            new_rId = new_slide.part.relate_to(rel.target_part, rel.reltype)
        rid_map[old_rId] = new_rId

    for shp in src.shapes:
        new_el = copy.deepcopy(shp._element)
        for el in new_el.iter():
            for attr_name, attr_val in list(el.attrib.items()):
                if attr_name.startswith(_R_NS) and attr_val in rid_map:
                    el.set(attr_name, rid_map[attr_val])
        new_slide.shapes._spTree.append(new_el)

    return new_slide


def move_slide(prs: Presentation, old_index: int, new_index: int):
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    if not 0 <= old_index < len(slides):
        return
    if not 0 <= new_index < len(slides):
        new_index = len(slides) - 1
    item = slides.pop(old_index)
    slides.insert(new_index, item)
    xml_slides.clear()
    for s in slides:
        xml_slides.append(s)


def delete_slide(prs: Presentation, index: int):
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    if not 0 <= index < len(slides):
        return
    target = slides[index]
    xml_slides.remove(target)
    prs.part.drop_rel(target.rId)


def find_shape_by_text(slide, contains: str):
    for shp in slide.shapes:
        try:
            if shp.has_text_frame and contains in shp.text_frame.text:
                return shp
        except Exception:
            continue
    return None


def set_text_preserve_format(shape, new_text: str):
    try:
        tf = shape.text_frame
        p = tf.paragraphs[0]
        if p.runs:
            p.runs[0].text = new_text
            for r in p.runs[1:]:
                r.text = ""
        else:
            p.text = new_text
    except Exception:
        try:
            shape.text = new_text
        except Exception:
            pass


def locate_question_template_slide(prs: Presentation) -> int:
    for i, slide in enumerate(prs.slides):
        if find_shape_by_text(slide, "#QUESTION") is not None:
            return i
    raise ValueError(
        "Sample PPT does not contain a slide with '#QUESTION'. "
        "Please add a slide with '#QUESTION' and 'Ans. (?)' placeholders."
    )


def add_question_image_to_slide(slide, img: Image.Image, content_left: Emu, content_top: Emu,
                                content_w: Emu, content_h: Emu):
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)

    iw, ih = img.size
    final_w, final_h = calculate_crop_size(iw, ih, int(content_w), int(content_h), mode="contain")
    left = int(content_left) + int((int(content_w) - final_w) / 2)
    top = int(content_top) + int((int(content_h) - final_h) / 2)

    # Remove any stale shapes in the content area before adding the new image.
    remove_overlapping_shapes(slide, left, top, left + final_w, top + final_h)
    slide.shapes.add_picture(buf, left, top, width=final_w, height=final_h)


def build_subject_ppt(template_bytes: bytes, questions: list, answers: dict, subject: str) -> bytes:
    """questions: list of dicts {number, image (PIL)} sorted by number."""
    prs = Presentation(io.BytesIO(template_bytes))
    template_idx = locate_question_template_slide(prs)

    slide_w, slide_h = prs.slide_width, prs.slide_height
    content_top = Emu(int(1.05 * 914400))
    content_left = Emu(int(0.45 * 914400))
    content_w = slide_w - Emu(int(0.9 * 914400))
    content_h = slide_h - content_top - Emu(int(0.35 * 914400))

    insert_at = template_idx
    for q in questions:
        new_slide = duplicate_slide(prs, template_idx)
        move_slide(prs, len(prs.slides) - 1, insert_at)
        insert_at += 1

        qshape = find_shape_by_text(new_slide, "#QUESTION")
        if qshape is not None:
            set_text_preserve_format(qshape, f"Q{q['number']}")

        ashape = find_shape_by_text(new_slide, "Ans. (")
        if ashape is not None:
            ans = answers.get(q["number"])
            set_text_preserve_format(ashape, f"Ans. ({ans})" if ans else "Ans. (—)")

        img = q["image"]
        if img is None:
            continue
        add_question_image_to_slide(new_slide, img, content_left, content_top, content_w, content_h)

    # Remove the original template slide after all new question slides are inserted.
    if insert_at < len(prs.slides):
        delete_slide(prs, insert_at)

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


# ---------------------------------------------------------------------------
# Main pipeline
# ---------------------------------------------------------------------------
def run_pipeline(pdf_bytes, template_bytes, ans_dict, client, model_name, progress_cb=None,
                 render_dpi_value=None, pad_px_value=None, col_gap_pct_value=None,
                 outer_margin_pct_value=None, invert_images=False):
    if render_dpi_value is not None:
        render_dpi_local = render_dpi_value
    else:
        render_dpi_local = render_dpi

    if pad_px_value is not None:
        pad_px_local = pad_px_value
    else:
        pad_px_local = pad_px

    if col_gap_pct_value is not None:
        col_gap_pct_local = col_gap_pct_value
    else:
        col_gap_pct_local = col_gap_pct

    if outer_margin_pct_value is not None:
        outer_margin_pct_local = outer_margin_pct_value
    else:
        outer_margin_pct_local = outer_margin_pct

    pages = convert_from_bytes(pdf_bytes, dpi=render_dpi_local)
    if not pages:
        raise ValueError("PDF rendering failed or no pages were found in the uploaded PDF.")

    all_questions: list[DetectedItem] = []
    all_headers: list[DetectedItem] = []

    total_steps = max(1, len(pages) * 2)
    current_step = 0

    for page_idx, page_img in enumerate(pages):
        columns = split_columns(page_img, outer_margin_pct_local, col_gap_pct_local)
        for col_idx, (col_img, off_x, off_y) in enumerate(columns):
            current_step += 1
            if progress_cb:
                progress_cb(current_step / total_steps, f"Page {page_idx + 1}/{len(pages)} — column {col_idx + 1}")

            raw_items = detect_items_in_column(client, model_name, col_img, KNOWN_SUBJECTS)
            if not isinstance(raw_items, list):
                continue

            cw, ch = col_img.size
            for it in raw_items:
                try:
                    if not isinstance(it, dict):
                        continue
                    box = it.get("box_2d")
                    if not isinstance(box, (list, tuple)) or len(box) != 4:
                        continue
                    x0, y0, x1, y1 = box_2d_to_abs(box, cw, ch, off_x, off_y)
                except Exception:
                    continue

                item_type = it.get("type")
                confidence = float(it.get("confidence", 0.8) or 0.8)

                if item_type == "question" and it.get("question_number") is not None:
                    try:
                        qnum = int(it["question_number"])
                    except Exception:
                        continue
                    all_questions.append(
                        DetectedItem(
                            kind="question",
                            page_index=page_idx,
                            column_index=col_idx,
                            y0=y0,
                            x0=x0,
                            y1=y1,
                            x1=x1,
                            question_number=qnum,
                            confidence=confidence,
                        )
                    )
                elif item_type == "section_header":
                    subject = normalize_subject(it.get("subject"), KNOWN_SUBJECTS)
                    if not subject:
                        continue
                    all_headers.append(
                        DetectedItem(
                            kind="section_header",
                            page_index=page_idx,
                            column_index=col_idx,
                            y0=y0,
                            x0=x0,
                            y1=y1,
                            x1=x1,
                            subject=subject,
                            confidence=confidence,
                        )
                    )

    if not all_questions:
        raise ValueError(
            "No question was detected. Please check the Gemini API key and the PDF/scan quality. "
            "You may also need to adjust the crop settings in the sidebar."
        )

    all_questions = deduplicate_questions(all_questions)

    by_col = {}
    for q in all_questions:
        by_col.setdefault((q.page_index, q.column_index), []).append(q)
    for items in by_col.values():
        clip_against_neighbours(items)

    all_headers.sort(key=lambda h: h.sort_key)
    all_questions.sort(key=lambda q: q.sort_key)

    current_subject = KNOWN_SUBJECTS[0]
    hi = 0
    for q in all_questions:
        while hi < len(all_headers) and all_headers[hi].sort_key <= q.sort_key:
            if all_headers[hi].subject:
                current_subject = all_headers[hi].subject
            hi += 1
        q.subject = current_subject

    # Crop and trim question images.
    subject_questions = {s: [] for s in KNOWN_SUBJECTS + ["Unclassified"]}
    for q in all_questions:
        page_img = pages[q.page_index]
        x0 = max(0, q.x0)
        y0 = max(0, q.y0)
        x1 = min(page_img.width, q.x1)
        y1 = min(page_img.height, q.y1)

        if x1 <= x0 or y1 <= y0:
            continue

        raw_crop = page_img.crop((x0, y0, x1, y1))
        trimmed = autotrim(raw_crop, pad=pad_px_local)

        try:
            validate_crop(trimmed, q.question_number)
        except Exception:
            continue

        final_img = invert_to_black_bg(trimmed) if invert_images else trimmed.convert("RGB")

        subject_name = normalize_subject(q.subject, KNOWN_SUBJECTS)
        if subject_name is None:
            subject_name = "Unclassified"

        subject_questions.setdefault(subject_name, []).append(
            {"number": q.question_number, "image": final_img, "page_index": q.page_index}
        )

    for subj in subject_questions:
        subject_questions[subj].sort(key=lambda d: d["number"])

    # Build output folder.
    work_dir = "output_processing"
    if os.path.exists(work_dir):
        shutil.rmtree(work_dir)
    os.makedirs(work_dir, exist_ok=True)

    subject_counts = {}
    for subj, qs in subject_questions.items():
        subject_counts[subj] = len(qs)
        if not qs:
            continue
        sub_folder = os.path.join(work_dir, subj)
        os.makedirs(sub_folder, exist_ok=True)
        ppt_bytes = build_subject_ppt(template_bytes, qs, ans_dict, subj)
        ppt_path = os.path.join(sub_folder, f"{subj}_Discussion.pptx")
        with open(ppt_path, "wb") as f:
            f.write(ppt_bytes)

    zip_path = "All_Subject_PPTs.zip"
    if os.path.exists(zip_path):
        os.remove(zip_path)
    shutil.make_archive("All_Subject_PPTs", "zip", work_dir)
    return zip_path, subject_counts, subject_questions


# ---------------------------------------------------------------------------
# UI
# ---------------------------------------------------------------------------
col1, col2 = st.columns(2)
with col1:
    pdf_file = st.file_uploader("1. Question Paper (PDF)", type=["pdf"])
with col2:
    template_ppt = st.file_uploader("2. Sample PPT Template (.pptx)", type=["pptx"])


tab1, tab2 = st.tabs(["Paste Answer Key", "Upload Answer Key Image/PDF (optional, OCR'd)"])
with tab1:
    ans_key_text = st.text_area(
        "Answer key",
        placeholder="Example:\n1: (3)\n2: (1)\n3: (4) ...",
        height=120,
    )
with tab2:
    ans_key_file = st.file_uploader("Answer key image/PDF", type=["png", "jpg", "jpeg", "pdf"], key="akf")


if st.button("🚀 Process & Generate All Subject PPTs", type="primary"):
    if not pdf_file or not template_ppt:
        st.error("⚠️ Please upload both the Question Paper PDF and the Sample PPT template.")
    elif not gemini_api_key:
        st.error("⚠️ Please add your Gemini API key in the sidebar.")
    else:
        try:
            client = get_client(gemini_api_key)

            ans_dict = parse_answer_key_text(ans_key_text)
            if ans_key_file is not None:
                if ans_key_file.type == "application/pdf":
                    ak_pages = convert_from_bytes(ans_key_file.read(), dpi=250)
                    for p in ak_pages:
                        ans_dict.update(ocr_answer_key_image(client, model_name, p))
                else:
                    img = Image.open(ans_key_file)
                    ans_dict.update(ocr_answer_key_image(client, model_name, img))

            progress_bar = st.progress(0.0)
            status_text = st.empty()

            def progress_cb(frac, msg):
                progress_bar.progress(min(frac, 1.0))
                status_text.text(msg)

            zip_path, subject_counts, subject_questions = run_pipeline(
                pdf_file.read(),
                template_ppt.read(),
                ans_dict,
                client,
                model_name,
                progress_cb,
                render_dpi_value=render_dpi,
                pad_px_value=pad_px,
                col_gap_pct_value=col_gap_pct,
                outer_margin_pct_value=outer_margin_pct,
                invert_images=invert_question_images,
            )

            progress_bar.empty()
            status_text.empty()
            st.success("🎉 All subject PPTs have been generated successfully!")

            st.write("**Subject-wise question count:**")
            st.table({"Subject": list(subject_counts.keys()), "Questions": list(subject_counts.values())})

            with st.expander("Preview a few cropped slides"):
                for subj, qs in subject_questions.items():
                    if not qs:
                        continue
                    st.markdown(f"**{subj}**")
                    preview_cols = st.columns(min(3, len(qs)))
                    for i, q in enumerate(qs[:3]):
                        with preview_cols[i]:
                            st.image(q["image"], caption=f"Q{q['number']}", use_container_width=True)

            with open(zip_path, "rb") as fp:
                st.download_button(
                    label="📦 Download All PPTs (ZIP)",
                    data=fp,
                    file_name=zip_path,
                    mime="application/zip",
                )

        except Exception as e:
            st.error(f"Error: {str(e)}")
            st.exception(e)
