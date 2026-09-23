"""Create self-contained demo assets for local verification."""
import io
from pathlib import Path

import fitz
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
(ROOT / "examples").mkdir(exist_ok=True)
(ROOT / "templates").mkdir(exist_ok=True)


def create_pdf():
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((48, 60), "JEE MAIN • DEMO QUESTION PAPER", fontsize=16)
    questions = [
        (110, "1. A particle moves with velocity v = 2t. Find its acceleration at t = 3 s."),
        (175, "2. If f(x) = x² + 2x, find f(2)."),
        (240, "3. Which of the following is an alkane?"),
        (305, "4. Find the value of 2 + 3 × 4."),
    ]
    for y, text in questions:
        page.insert_text((55, y), text, fontsize=11)
        page.insert_text((75, y + 28), "(A) Option A     (B) Option B     (C) Option C     (D) Option D", fontsize=9)
    data = doc.tobytes()
    doc.close()
    (ROOT / "examples" / "demo_question_paper.pdf").write_bytes(data)
    (ROOT / "examples" / "demo_answer_key.txt").write_text("1: A\n2: B\n3: C\n4: D\n", encoding="utf-8")


def create_template():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = RGBColor(11, 15, 22)

    header = slide.shapes.add_textbox(Inches(.55), Inches(.28), Inches(8.5), Inches(.55))
    p = header.text_frame.paragraphs[0]
    p.text = "VIDYAPEETH  •  TEST DISCUSSION"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = RGBColor(224, 238, 255)

    meta = slide.shapes.add_textbox(Inches(9.0), Inches(.28), Inches(3.7), Inches(.55))
    p = meta.text_frame.paragraphs[0]
    p.text = "#SUBJECT   |   #QUESTION"
    p.alignment = PP_ALIGN.RIGHT
    p.font.size = Pt(14)
    p.font.color.rgb = RGBColor(93, 210, 255)

    accent = slide.shapes.add_shape(1, Inches(.55), Inches(.92), Inches(12.2), Inches(.025))
    accent.fill.solid(); accent.fill.fore_color.rgb = RGBColor(93, 210, 255); accent.line.fill.background()

    area = slide.shapes.add_shape(1, Inches(.75), Inches(1.25), Inches(11.85), Inches(5.25))
    area.name = "QUESTION_IMAGE"
    area.fill.solid(); area.fill.fore_color.rgb = RGBColor(245, 247, 250)
    area.line.color.rgb = RGBColor(45, 57, 73)

    footer = slide.shapes.add_textbox(Inches(.75), Inches(6.75), Inches(11.85), Inches(.35))
    p = footer.text_frame.paragraphs[0]
    p.text = "ANSWER  #ANSWER   •   Local-first Presentation Studio"
    p.font.size = Pt(10); p.font.color.rgb = RGBColor(157, 171, 190)

    output = io.BytesIO(); prs.save(output)
    (ROOT / "templates" / "Vidyapeeth_Premium_Discussion_Template.pptx").write_bytes(output.getvalue())


if __name__ == "__main__":
    create_pdf(); create_template(); print("Demo assets created.")
