from pptx import Presentation
from pptx.util import Inches


def build_subject_ppt(template_path, image_paths, output_ppt_path):
    prs = Presentation(template_path)
    blank_layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0]

    for img_path in image_paths:
        slide = prs.slides.add_slide(blank_layout)
        left = Inches(0.8)
        top = Inches(1.2)
        width = Inches(8.4)
        slide.shapes.add_picture(img_path, left, top, width=width)

    prs.save(output_ppt_path)