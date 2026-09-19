import copy
import io
from pptx import Presentation
from pptx.util import Emu

# ... keep existing duplicate_slide, move_slide, delete_slide, find_shape_by_text, set_shape_text ...
# (Insert the utility functions from your original ppt_service.py here)

def build_subject_ppt(template_bytes: bytes, questions: list[dict], answers: dict[int, str]) -> bytes:
    prs = Presentation(io.BytesIO(template_bytes))
    
    # Locate the template slide containing the #QUESTION placeholder
    template_index = next((i for i, slide in enumerate(prs.slides) if find_shape_by_text(slide, "#QUESTION")), None)
    if template_index is None:
        raise ValueError("The PPT template must contain a #QUESTION placeholder.")

    slide_width, slide_height = prs.slide_width, prs.slide_height
    # Adjust content coordinates to fit beneath the Vidyapeeth headers
    content_left, content_top = Emu(int(0.45 * 914400)), Emu(int(1.50 * 914400))
    content_width = slide_width - Emu(int(0.9 * 914400))
    content_height = slide_height - content_top - Emu(int(0.35 * 914400))
    
    insert_index = template_index

    for question in questions:
        slide = duplicate_slide(prs, template_index)
        move_slide(prs, len(prs.slides) - 1, insert_index)
        insert_index += 1

        if q_shape := find_shape_by_text(slide, "#QUESTION"):
            set_shape_text(q_shape, f"Q{question['number']}")

        # Matches the 'Ans. (?)' format from the provided PPT template
        if ans_shape := find_shape_by_text(slide, "Ans. ("):
            ans_val = answers.get(question['number'], "?")
            set_shape_text(ans_shape, f"Ans. ({ans_val})")

        image_buffer = io.BytesIO()
        question["image"].save(image_buffer, format="PNG")
        image_buffer.seek(0)

        img_w, img_h = question["image"].size
        scale = min(int(content_width) / img_w, int(content_height) / img_h)
        final_w, final_h = max(1, int(img_w * scale)), max(1, int(img_h * scale))
        left = int(content_left) + (int(content_width) - final_w) // 2
        top = int(content_top) + (int(content_height) - final_h) // 2

        remove_content_shapes(slide, left, top, left + final_w, top + final_h)
        slide.shapes.add_picture(image_buffer, left, top, width=final_w, height=final_h)

    delete_slide(prs, insert_index)
    output = io.BytesIO()
    prs.save(output)
    return output.getvalue()
