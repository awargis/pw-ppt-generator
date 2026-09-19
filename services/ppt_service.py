import copy
import io

from pptx import Presentation
from pptx.util import Emu


RELATIONSHIP_NAMESPACE = (
    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
)


def find_shape_by_text(slide, text: str):
    for shape in slide.shapes:
        try:
            if shape.has_text_frame and text in shape.text_frame.text:
                return shape
        except Exception:
            continue

    return None


def set_shape_text(shape, text: str):
    text_frame = shape.text_frame
    paragraph = text_frame.paragraphs[0]

    if paragraph.runs:
        paragraph.runs[0].text = text

        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.text = text


def duplicate_slide(prs: Presentation, source_index: int):
    source = prs.slides[source_index]
    new_slide = prs.slides.add_slide(source.slide_layout)

    for shape in list(new_slide.shapes):
        shape._element.getparent().remove(shape._element)

    relationship_map = {}

    for old_id, relationship in source.part.rels.items():
        if "notesSlide" in relationship.reltype:
            continue

        if "slideLayout" in relationship.reltype:
            continue

        if relationship.is_external:
            new_id = new_slide.part.relate_to(
                relationship.target_ref,
                relationship.reltype,
                is_external=True,
            )
        else:
            new_id = new_slide.part.relate_to(
                relationship.target_part,
                relationship.reltype,
            )

        relationship_map[old_id] = new_id

    for shape in source.shapes:
        copied_element = copy.deepcopy(shape._element)

        for element in copied_element.iter():
            for attribute, value in list(element.attrib.items()):
                if (
                    attribute.startswith(RELATIONSHIP_NAMESPACE)
                    and value in relationship_map
                ):
                    element.set(attribute, relationship_map[value])

        new_slide.shapes._spTree.append(copied_element)

    return new_slide


def delete_slide(prs: Presentation, index: int):
    slide_ids = prs.slides._sldIdLst
    slides = list(slide_ids)

    if index < 0 or index >= len(slides):
        return

    target = slides[index]
    slide_ids.remove(target)
    prs.part.drop_rel(target.rId)


def move_slide(prs: Presentation, old_index: int, new_index: int):
    slide_ids = prs.slides._sldIdLst
    slides = list(slide_ids)

    if old_index < 0 or old_index >= len(slides):
        return

    item = slides.pop(old_index)
    new_index = max(0, min(new_index, len(slides)))
    slides.insert(new_index, item)

    slide_ids.clear()

    for slide in slides:
        slide_ids.append(slide)


def remove_content_shapes(slide, left, top, right, bottom):
    for shape in list(slide.shapes):
        try:
            text = shape.text_frame.text if shape.has_text_frame else ""

            if "#QUESTION" in text or "Ans." in text:
                continue

            shape_right = shape.left + shape.width
            shape_bottom = shape.top + shape.height

            overlaps = (
                shape.left < right
                and shape_right > left
                and shape.top < bottom
                and shape_bottom > top
            )

            if overlaps:
                shape._element.getparent().remove(shape._element)

        except Exception:
            continue


def build_subject_ppt(
    template_bytes: bytes,
    questions: list[dict],
    answers: dict[int, str],
) -> bytes:
    prs = Presentation(io.BytesIO(template_bytes))

    template_index = None

    for index, slide in enumerate(prs.slides):
        if find_shape_by_text(slide, "#QUESTION"):
            template_index = index
            break

    if template_index is None:
        raise ValueError(
            "The PPT template must contain a #QUESTION placeholder."
        )

    slide_width = prs.slide_width
    slide_height = prs.slide_height

    content_left = Emu(int(0.45 * 914400))
    content_top = Emu(int(1.05 * 914400))
    content_width = slide_width - Emu(int(0.9 * 914400))
    content_height = slide_height - content_top - Emu(int(0.35 * 914400))

    insert_index = template_index

    for question in questions:
        slide = duplicate_slide(prs, template_index)
        move_slide(prs, len(prs.slides) - 1, insert_index)
        insert_index += 1

        question_shape = find_shape_by_text(slide, "#QUESTION")
        if question_shape:
            set_shape_text(question_shape, f"Q{question['number']}")

        answer_shape = find_shape_by_text(slide, "Ans. (")
        if answer_shape:
            answer = answers.get(question["number"], "—")
            set_shape_text(answer_shape, f"Ans. ({answer})")

        image = question["image"]
        image_buffer = io.BytesIO()
        image.save(image_buffer, format="PNG")
        image_buffer.seek(0)

        image_width, image_height = image.size
        scale = min(
            int(content_width) / image_width,
            int(content_height) / image_height,
        )

        final_width = max(1, int(image_width * scale))
        final_height = max(1, int(image_height * scale))

        left = int(content_left) + (
            int(content_width) - final_width
        ) // 2

        top = int(content_top) + (
            int(content_height) - final_height
        ) // 2

        remove_content_shapes(
            slide,
            left,
            top,
            left + final_width,
            top + final_height,
        )

        slide.shapes.add_picture(
            image_buffer,
            left,
            top,
            width=final_width,
            height=final_height,
        )

    delete_slide(prs, insert_index)

    output = io.BytesIO()
    prs.save(output)
    return output.getvalue()
