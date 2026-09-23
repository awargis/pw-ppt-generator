import io

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from models import BoundingBox, QuestionRegion
from ppt.exporter import export_subject_ppts


def make_template() -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.4), Inches(0.2), Inches(9), Inches(0.5))
    title.text_frame.text = "#QUESTION | #SUBJECT | Answer: #ANSWER"
    image_area = slide.shapes.add_textbox(Inches(0.8), Inches(1.0), Inches(8), Inches(0.3))
    image_area.name = "QUESTION_IMAGE"
    output = io.BytesIO()
    prs.save(output)
    return output.getvalue()


def test_template_based_ppt_generation_creates_nonempty_ppt():
    image = Image.new("RGB", (800, 500), "white")
    region = QuestionRegion(
        number=1, page_index=0, column_index=0,
        box=BoundingBox(0, 0, 800, 500), image=image, subject="Mathematics",
    )
    outputs = export_subject_ppts(make_template(), [region], {1: "A"})
    assert set(outputs) == {"Mathematics"}
    assert len(outputs["Mathematics"]) > 1000
    prs = Presentation(io.BytesIO(outputs["Mathematics"]))
    assert len(prs.slides) == 1
    assert any(shape.shape_type == 13 for shape in prs.slides[0].shapes)
    text = " ".join(shape.text for shape in prs.slides[0].shapes if shape.has_text_frame)
    assert "Q1" in text and "Mathematics" in text and "A" in text
